"""Tests for turning prices into knapsack items and writing the Excel report.

Run with:  pytest -q
"""

import openpyxl

from main import build_stocks
from simulator.config import COMPANIES
from simulator.optimization import dp_allocation, greedy_allocation
from simulator.reporting import save_to_excel


def test_build_stocks_applies_expected_returns():
    prices = {name: 1000.0 for name in COMPANIES}
    stocks = {name: (cost, value) for name, cost, value in build_stocks(prices)}
    assert set(stocks) == set(COMPANIES)
    for name, (_, return_pct) in COMPANIES.items():
        assert stocks[name] == (1000, int(1000 * (1 + return_pct / 100)))


def test_build_stocks_skips_missing_prices():
    prices = {"TCS": 3245.5, "HDFC": 0.0}
    names = [name for name, _, _ in build_stocks(prices)]
    assert names == ["TCS"]


def test_excel_report_has_all_sheets(tmp_path):
    stocks = [("A", 6, 9), ("B", 5, 7)]
    greedy = greedy_allocation(stocks, 10)
    dp = dp_allocation(stocks, 10)

    path = save_to_excel({"Mode": "test"}, greedy, dp, None, tmp_path / "report.xlsx")

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Prediction", "Portfolio", "Stock Quantities"]
    rows = list(wb["Stock Quantities"].iter_rows(values_only=True))
    assert rows == [("Company", "Greedy Quantity", "DP Quantity"), ("A", 1, 0), ("B", 0, 2)]
